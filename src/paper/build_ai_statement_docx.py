from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / ".deps"))

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = ROOT / "paper" / "AI工具使用说明.docx"


def set_run_font(run, east_asia="宋体", latin="Times New Roman", size=10.5, bold=False):
    run.font.name = latin
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:eastAsia"), east_asia)


def set_paragraph_format(paragraph, first_line=True, before=0, after=0):
    fmt = paragraph.paragraph_format
    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    fmt.line_spacing = 1.3
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.first_line_indent = Cm(0.74) if first_line else Cm(0)
    fmt.keep_together = True


def add_body(doc, text, first_line=True):
    paragraph = doc.add_paragraph()
    set_paragraph_format(paragraph, first_line=first_line)
    set_run_font(paragraph.add_run(text))
    return paragraph


def add_section_heading(doc, text):
    paragraph = doc.add_paragraph()
    set_paragraph_format(paragraph, first_line=False, before=6, after=2)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.keep_with_next = True
    set_run_font(paragraph.add_run(text), east_asia="黑体", size=10.5, bold=True)
    return paragraph


def set_cell_margins(cell, top="80", start="100", bottom="80", end="100"):
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), value)
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, edge, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    border = borders.find(qn(f"w:{edge}"))
    if border is None:
        border = OxmlElement(f"w:{edge}")
        borders.append(border)
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), size)
    border.set(qn("w:color"), color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")


def set_cell_text(cell, text, header=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    set_cell_margins(cell)
    for edge in ("top", "left", "bottom", "right"):
        set_cell_border(cell, edge)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.08
    run = paragraph.add_run(text)
    set_run_font(
        run,
        east_asia="黑体" if header else "宋体",
        size=9.5,
        bold=header,
    )
    if header:
        run.font.color.rgb = RGBColor(255, 255, 255)


def add_interaction_table(doc):
    columns = ["编号", "交互主题", "典型交互摘要", "采纳与核验"]
    records = [
        (
            "1",
            "题面解析与问题依赖",
            "依据题面和附件，要求梳理四问目标、决策变量、约束及相互依赖，比较问题类型。",
            "形成Q1机制-优化、Q2/Q3预测-优化、Q4情景分析-优化的表述；再对照题面与问题依赖图确认。",
        ),
        (
            "2",
            "数据口径与统一模型",
            "根据负荷、光伏、电价文件，要求检查时间粒度、单位、能量换算和储能边界。",
            "完善统一符号表与模型假设；通过数据报告、原始附件和程序读取结果核对口径。",
        ),
        (
            "3",
            "优化模型与程序辅助",
            "要求给出能量平衡、储能状态转移、购电费用目标及线性规划求解框架，并辅助调试。",
            "以程序运行结果为准；复核终端SOC、功率边界、能量平衡和网格收敛，不采纳未经运行验证的结果。",
        ),
        (
            "4",
            "预测与滚动调整",
            "要求比较非前视预测方案，分析光伏更新后的滚动调整，并讨论安全裕度参数的费用-供能缺口取舍。",
            "核查训练和决策时点，使用结果工作簿及稳健性报告验证；将0.99的选择限定为既定风险偏好下的结论。",
        ),
        (
            "5",
            "价格情景与论文表达",
            "要求区分历史价格情景鲁棒优化和真实未来价格参照，并据冻结数字生成图表、表格和文字初稿。",
            "检查未来价格泄漏、图表数字与正文一致性；保留历史情景覆盖有限和外推受限的说明。",
        ),
    ]
    table = doc.add_table(rows=1, cols=len(columns))
    table.autofit = False
    table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    widths = [Cm(1.0), Cm(2.7), Cm(6.0), Cm(6.3)]
    for index, width in enumerate(widths):
        table.columns[index].width = width
        for cell in table.columns[index].cells:
            cell.width = width
    for index, label in enumerate(columns):
        shade_cell(table.rows[0].cells[index], "1F4E78")
        set_cell_text(table.rows[0].cells[index], label, header=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    for row_index, record in enumerate(records):
        row = table.add_row()
        for column_index, value in enumerate(record):
            cell = row.cells[column_index]
            if row_index % 2 == 1:
                shade_cell(cell, "EAF2F8")
            alignment = WD_ALIGN_PARAGRAPH.CENTER if column_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            set_cell_text(cell, value, align=alignment)
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    return table


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)
    set_run_font(run, size=10.5)


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    title_style = doc.styles["Title"]
    title_style.font.name = "Times New Roman"
    title_style.font.size = Pt(16)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    title_style_ppr = title_style._element.get_or_add_pPr()
    title_border = title_style_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_style_ppr.remove(title_border)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    title.paragraph_format.keep_with_next = True
    title_ppr = title._p.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)
    set_run_font(title.add_run("AI工具使用说明"), east_asia="黑体", size=16, bold=True)

    add_body(
        doc,
        "本参赛队在竞赛过程中使用了人工智能工具，主要用于题意梳理、候选方法比较、程序辅助、结果核查、图表与文档排版以及语言润色。模型方案、参数选择、实验结果和论文表述均由参赛队结合题目附件、程序运行记录与结果文件审查后确认，参赛队对最终提交内容承担责任。",
    )

    add_section_heading(doc, "一  使用的AI工具")
    add_body(
        doc,
        "使用的工具为OpenAI ChatGPT与Codex，使用平台在线提供的GPT-5系列模型及其代码辅助能力。由于在线服务会持续更新，具体子版本以相应会话和平台记录为准。",
    )

    add_section_heading(doc, "二  使用目的与环节")
    add_body(
        doc,
        "AI主要参与以下辅助环节：第一，整理四个问题之间的依赖关系，比较确定性优化、非前视预测优化、滚动调整和历史价格情景鲁棒优化等候选方案；第二，辅助编写和调试Python程序，检查单位换算、储能状态转移、能量平衡、非前视边界与输出格式；第三，协助汇总实验结果，生成图表初稿并检查论文中的符号、数值和文件名称是否一致；第四，对论文草稿进行语言润色和Word排版。AI未替代原始数据、实际程序运行结果或参赛队对最终方案的确认。",
    )

    add_section_heading(doc, "三  主要提示方式与使用过程")
    add_body(
        doc,
        "使用过程中，参赛队按任务逐步提供题面、附件数据、已确定的建模口径和阶段性结果，并通过限定输入、输出和验证条件约束AI。典型提示包括：解读题面并梳理各问的目标与约束；比较候选模型的适用性、可解释性和竞赛实现成本；根据给定数据生成可复现的程序框架；检查未来信息是否进入当前决策；依据已经冻结的实验数字撰写结果分析；核对图表、正文与结果文件之间的一致性。对于影响建模结论的内容，采用“提出候选方案—运行程序—核对结果—修改表述”的方式迭代处理。",
    )

    doc.add_page_break()
    add_section_heading(doc, "四  典型AI使用交互记录")
    add_body(
        doc,
        "下表为实际使用环节的摘要性记录，不替代完整会话日志；其中“典型交互摘要”概括任务输入，“采纳与核验”说明输出进入作品前的处理方式。",
    )
    add_interaction_table(doc)

    add_section_heading(doc, "五  采纳 修改与核验情况")
    add_body(
        doc,
        "AI输出主要作为分析建议、代码草稿和文字初稿使用。参赛队依据题目附件修正变量定义、时间尺度和结算规则，并通过程序运行、单元测试和结果工作簿核对目标函数值、储能边界、能量平衡和关键统计量；同时检查非前视信息边界、图表正文一致性及适用范围。语言润色类建议仅在不改变技术含义的前提下采用，无法由数据或计算结果支持的表述未作为结论使用。",
    )

    add_body(
        doc,
        "本说明依据《全国大学生数学建模竞赛人工智能工具使用规定（2026年试行）》编制，用于公开披露本作品中AI工具的使用情况。",
    )

    add_page_number(section.footer.paragraphs[0])
    core = doc.core_properties
    core.title = "AI工具使用说明"
    core.subject = "数学建模竞赛论文AI工具使用披露"
    core.author = "参赛队"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
