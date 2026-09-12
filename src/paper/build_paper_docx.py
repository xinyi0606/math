"""Build the contest paper DOCX from the canonical Markdown draft.

Equations are intentionally preserved as plain LaTeX text because the current
revision explicitly excludes formula-layout conversion.
"""

from __future__ import annotations

import re
import sys
from pathlib import Path


ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / ".deps"))

from docx import Document  # noqa: E402
from docx.enum.section import WD_SECTION  # noqa: E402
from docx.enum.style import WD_STYLE_TYPE  # noqa: E402
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402


SOURCE = ROOT / "paper" / "sections" / "complete_first_draft.md"
OUTPUT = ROOT / "paper" / "光储微电网多时间尺度协调调度论文正文.docx"


def set_run_font(run, east_asia: str = "宋体", latin: str = "Times New Roman", size: float = 10.5, bold=None):
    run.font.name = latin
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, text, end):
        run._r.append(element)
    set_run_font(run, size=9)


def clean_inline(text: str) -> str:
    text = re.sub(r"<!--.*?-->", "", text).strip()
    text = text.replace("**", "")
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1（\2）", text)
    return text


def add_text_paragraph(doc: Document, text: str, *, center: bool = False, indent: bool = True, size: float = 10.5):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    paragraph.paragraph_format.space_after = Pt(2)
    if indent and not center:
        paragraph.paragraph_format.first_line_indent = Pt(21)
    run = paragraph.add_run(clean_inline(text))
    set_run_font(run, size=size)
    return paragraph


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    add_page_number(section.footer.paragraphs[0])

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE

    styles = {
        "Title": ("黑体", 18, True, WD_ALIGN_PARAGRAPH.CENTER),
        "Heading 1": ("黑体", 14, True, WD_ALIGN_PARAGRAPH.LEFT),
        "Heading 2": ("黑体", 12, True, WD_ALIGN_PARAGRAPH.LEFT),
    }
    for name, (font_name, size, bold, alignment) in styles.items():
        style = doc.styles[name]
        style.font.name = font_name
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
        style.paragraph_format.alignment = alignment
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(5)
        if name == "Title":
            style_properties = style._element.get_or_add_pPr()
            borders = style_properties.find(qn("w:pBdr"))
            if borders is not None:
                style_properties.remove(borders)

    if "Figure Caption" not in [style.name for style in doc.styles]:
        style = doc.styles.add_style("Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = doc.styles["Figure Caption"]
    style.font.name = "宋体"
    style.font.size = Pt(9)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style.paragraph_format.space_after = Pt(5)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    table.autofit = True
    for index, value in enumerate(rows[0]):
        cell = table.rows[0].cells[index]
        cell.text = clean_inline(value)
        set_cell_shading(cell, "D9EAF7")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            set_run_font(run, east_asia="黑体", size=8.5, bold=True)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    repeat_table_header(table.rows[0])
    for values in rows[1:]:
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cells[index].text = clean_inline(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cells[index].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in cells[index].paragraphs[0].runs:
                set_run_font(run, size=8.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build() -> None:
    doc = Document()
    configure_document(doc)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    figure_number = 0
    i = 0
    while i < len(lines):
        raw = lines[i].rstrip()
        line = raw.strip()
        if not line:
            i += 1
            continue
        if line.startswith("# "):
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.space_after = Pt(5)
            paragraph.paragraph_format.keep_with_next = True
            run = paragraph.add_run(clean_inline(line[2:]))
            set_run_font(run, east_asia="黑体", latin="Times New Roman", size=18, bold=True)
            i += 1
            continue
        if line.startswith("## "):
            paragraph = doc.add_paragraph(style="Heading 1")
            run = paragraph.add_run(clean_inline(line[3:]))
            set_run_font(run, east_asia="黑体", size=14, bold=True)
            i += 1
            continue
        if line.startswith("### "):
            paragraph = doc.add_paragraph(style="Heading 2")
            run = paragraph.add_run(clean_inline(line[4:]))
            set_run_font(run, east_asia="黑体", size=12, bold=True)
            i += 1
            continue
        image_match = re.fullmatch(r"!\[([^\]]+)\]\(([^)]+)\)", line)
        if image_match:
            figure_number += 1
            caption, relative = image_match.groups()
            image_path = (SOURCE.parent / relative).resolve()
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run()
            run.add_picture(str(image_path), width=Cm(15.4))
            caption_paragraph = doc.add_paragraph(style="Figure Caption")
            caption_run = caption_paragraph.add_run(f"图{figure_number}  {caption}")
            set_run_font(caption_run, size=9)
            i += 1
            continue
        if line.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            parsed = [[cell.strip() for cell in row.strip("|").split("|")] for row in table_lines]
            if len(parsed) >= 2 and all(re.fullmatch(r":?-+:?", cell) for cell in parsed[1]):
                parsed.pop(1)
            add_table(doc, parsed)
            continue
        if line.startswith("$$"):
            formula_lines = [line]
            if line.count("$$") < 2:
                i += 1
                while i < len(lines):
                    formula_lines.append(lines[i].strip())
                    if "$$" in lines[i]:
                        break
                    i += 1
            add_text_paragraph(doc, " ".join(formula_lines), center=True, indent=False, size=10)
            i += 1
            continue
        if re.match(r"^\d+\.\s", line):
            paragraph = add_text_paragraph(doc, line, indent=False)
            paragraph.paragraph_format.left_indent = Pt(21)
            paragraph.paragraph_format.first_line_indent = Pt(-21)
            i += 1
            continue
        if line.startswith(">"):
            paragraph = add_text_paragraph(doc, line.lstrip("> "), indent=False, size=9)
            paragraph.paragraph_format.left_indent = Pt(21)
            i += 1
            continue
        add_text_paragraph(doc, line)
        i += 1

    doc.core_properties.title = "面向预测不确定性与波动电价的光储微电网多时间尺度协调调度"
    doc.core_properties.subject = "数学建模竞赛论文"
    doc.core_properties.author = ""
    temporary = OUTPUT.with_suffix(".tmp.docx")
    doc.save(temporary)
    temporary.replace(OUTPUT)


if __name__ == "__main__":
    build()
