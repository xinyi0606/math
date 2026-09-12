from __future__ import annotations

import json
import shutil
from datetime import datetime
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook

from .config import DATA_TEMPLATES, OUTPUT, RESULTS, TOLERANCE


def round_directories(qx: str, round_number: int = 1) -> dict[str, Path]:
    root = RESULTS / qx / "experiments" / f"round{round_number}"
    paths = {name: root / name for name in ("figures", "tables", "metrics", "logs")}
    paths["root"] = root
    for path in paths.values():
        path.mkdir(parents=True, exist_ok=True)
    OUTPUT.mkdir(parents=True, exist_ok=True)
    return paths


def _json_default(value: Any):
    if isinstance(value, (np.floating, np.integer)):
        return value.item()
    if isinstance(value, np.ndarray):
        return value.tolist()
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, (datetime, np.datetime64)):
        return str(value)
    raise TypeError(f"Cannot serialize {type(value)}")


def write_json(path: Path, payload: dict[str, Any]) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2, default=_json_default) + "\n", encoding="utf-8")


def copy_official_result(source: Path, official_name: str) -> Path:
    target = OUTPUT / official_name
    shutil.copy2(source, target)
    return target


def write_q1_workbook(path: Path, grid: np.ndarray, charge: np.ndarray, discharge: np.ndarray, soc: np.ndarray) -> None:
    workbook = load_workbook(DATA_TEMPLATES / "result1.xlsx")
    plan = workbook["计划购电量"]
    for index, value in enumerate(grid, start=2):
        plan.cell(index, 2, float(value))
    storage = workbook["充放电量"]
    for block in range(6):
        row = block + 2
        left, right = block * 24, (block + 1) * 24
        storage.cell(row, 2, float(charge[left:right].sum()))
        storage.cell(row, 3, float(discharge[left:right].sum()))
    storage.cell(2, 5, float(soc[0]))
    storage.cell(3, 5, float(soc[-1]))
    workbook.save(path)


def _date_key(value) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, np.datetime64):
        return str(value.astype("datetime64[D]"))
    return None


def _nonnegative_output(value: float) -> float:
    number = float(value)
    if not np.isfinite(number):
        raise ValueError(f"Official workbook value is not finite: {number}")
    if number < -TOLERANCE:
        raise ValueError(f"Official workbook value is materially negative: {number}")
    return max(number, 0.0)


def _fill_daily_matrix(sheet, dates: np.ndarray, values: np.ndarray, daily_cost: np.ndarray) -> None:
    date_lookup = {str(date.astype("datetime64[D]")): index for index, date in enumerate(dates)}
    for row in range(2, sheet.max_row + 1):
        key = _date_key(sheet.cell(row, 1).value)
        if key not in date_lookup:
            continue
        index = date_lookup[key]
        normalized = np.array([_nonnegative_output(value) for value in values[index]])
        for slot, value in enumerate(normalized, start=2):
            sheet.cell(row, slot, float(value))
        sheet.cell(row, 146, float(normalized.sum()))
        sheet.cell(row, 147, float(daily_cost[index]))


def _fill_storage_samples(sheet, dates: np.ndarray, charge: np.ndarray, discharge: np.ndarray, soc: np.ndarray) -> None:
    date_lookup = {str(date.astype("datetime64[D]")): index for index, date in enumerate(dates)}
    current_key = None
    for row in range(2, sheet.max_row + 1):
        candidate = _date_key(sheet.cell(row, 1).value)
        if candidate is not None:
            current_key = candidate
        if current_key not in date_lookup:
            continue
        label = sheet.cell(row, 2).value
        if not isinstance(label, str) or "-" not in label:
            continue
        index = date_lookup[current_key]
        block = {"0:00-4:00": 0, "4:00-8:00": 1, "8:00-12:00": 2,
                 "12:00-16:00": 3, "16:00-20:00": 4, "20:00-24:00": 5}.get(label)
        if block is None:
            continue
        left, right = 24 * block, 24 * (block + 1)
        sheet.cell(row, 3, float(charge[index, left:right].sum()))
        sheet.cell(row, 4, float(discharge[index, left:right].sum()))
        moment = sheet.cell(row, 5).value
        if str(moment) in ("00:00:00", "0:00", "00:00"):
            sheet.cell(row, 6, float(soc[index, 0]))
        elif str(moment) == "24:00":
            sheet.cell(row, 6, float(soc[index, -1]))


def contiguous_emergency_events(values: np.ndarray, labels: list[str], tolerance: float = 1e-8):
    active = np.asarray(values) > tolerance
    events = []
    start = None
    for index, flag in enumerate(np.r_[active, False]):
        if flag and start is None:
            start = index
        elif not flag and start is not None:
            end = index
            label = labels[start] if end - start == 1 else f"{labels[start]}—{labels[end - 1]}"
            events.append((label, float(np.asarray(values)[start:end].sum())))
            start = None
    return events


def _fill_emergency_samples(sheet, dates: np.ndarray, emergency: np.ndarray, labels: list[str]) -> None:
    date_lookup = {str(date.astype("datetime64[D]")): index for index, date in enumerate(dates)}
    current_key = None
    group_row = 0
    selected_events = []
    for row in range(2, sheet.max_row + 1):
        candidate = _date_key(sheet.cell(row, 1).value)
        if candidate is not None:
            current_key = candidate
            group_row = 0
            if current_key in date_lookup:
                events = contiguous_emergency_events(emergency[date_lookup[current_key]], labels)
                selected_events = sorted(sorted(events, key=lambda item: item[1], reverse=True)[:3])
            else:
                selected_events = []
        if current_key in date_lookup and group_row < len(selected_events):
            label, amount = selected_events[group_row]
            sheet.cell(row, 2, label)
            sheet.cell(row, 3, amount)
        group_row += 1


def write_q2_workbook(
    path: Path,
    dates: np.ndarray,
    plan_grid: np.ndarray,
    plan_cost: np.ndarray,
    charge: np.ndarray,
    discharge: np.ndarray,
    soc: np.ndarray,
    emergency: np.ndarray,
    template_name: str = "result2.xlsx",
) -> None:
    workbook = load_workbook(DATA_TEMPLATES / template_name)
    _fill_daily_matrix(workbook["计划购电量"], dates, plan_grid, plan_cost)
    _fill_storage_samples(workbook["充放电量"], dates, charge, discharge, soc)
    labels = [workbook["计划购电量"].cell(1, column).value for column in range(2, 146)]
    _fill_emergency_samples(workbook["紧急购电量"], dates, emergency, labels)
    workbook.save(path)


def write_q3_workbook(
    path: Path,
    dates: np.ndarray,
    plan_grid: np.ndarray,
    plan_cost: np.ndarray,
    adjusted_grid: np.ndarray,
    adjusted_contract_cost: np.ndarray,
    charge: np.ndarray,
    discharge: np.ndarray,
    soc: np.ndarray,
    emergency: np.ndarray,
    template_name: str = "result3.xlsx",
) -> None:
    workbook = load_workbook(DATA_TEMPLATES / template_name)
    _fill_daily_matrix(workbook["计划购电量"], dates, plan_grid, plan_cost)
    _fill_daily_matrix(workbook["调整购电量"], dates, adjusted_grid, adjusted_contract_cost)
    _fill_storage_samples(workbook["充放电量"], dates, charge, discharge, soc)
    labels = [workbook["计划购电量"].cell(1, column).value for column in range(2, 146)]
    _fill_emergency_samples(workbook["紧急购电量"], dates, emergency, labels)
    workbook.save(path)
