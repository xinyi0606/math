import hashlib
import json
import unittest
from pathlib import Path

import math
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]


class FormalOutputTest(unittest.TestCase):
    def test_original_inputs_and_templates_keep_manifest_hashes(self):
        for line in (ROOT / "data" / "manifest.sha256").read_text(encoding="utf-8").splitlines():
            expected, relative = line.split(maxsplit=1)
            payload = (ROOT / "data" / relative).read_bytes()
            self.assertEqual(hashlib.sha256(payload).hexdigest(), expected, relative)

    def test_official_workbooks_preserve_template_structure_and_size(self):
        for name in ("result1.xlsx", "result2.xlsx", "result3.xlsx", "result4-2.xlsx", "result4-3.xlsx"):
            with self.subTest(name=name):
                template = load_workbook(ROOT / "data" / "templates" / name, read_only=True)
                result_path = ROOT / "output" / name
                result = load_workbook(result_path, read_only=True, data_only=True)
                self.assertEqual(result.sheetnames, template.sheetnames)
                self.assertEqual([(w.max_row, w.max_column) for w in result.worksheets],
                                 [(w.max_row, w.max_column) for w in template.worksheets])
                self.assertLess(result_path.stat().st_size, 5 * 1024 * 1024)

    def test_all_plan_and_adjusted_matrices_are_finite(self):
        for name in ("result2.xlsx", "result3.xlsx", "result4-2.xlsx", "result4-3.xlsx"):
            workbook = load_workbook(ROOT / "output" / name, read_only=True, data_only=True)
            sheets = ["计划购电量"] + (["调整购电量"] if "调整购电量" in workbook.sheetnames else [])
            for sheet_name in sheets:
                sheet = workbook[sheet_name]
                for row in sheet.iter_rows(min_row=2, min_col=2, max_col=145, values_only=True):
                    self.assertTrue(all(value is not None and math.isfinite(value) and value >= 0 for value in row),
                                    f"{name}/{sheet_name}")

    def test_run_summaries_and_selected_constraint_audits_pass(self):
        for qx in ("Q1", "Q2", "Q3", "Q4"):
            summary_path = ROOT / "results" / qx / "experiments" / "round1" / "run_summary.json"
            summary = json.loads(summary_path.read_text(encoding="utf-8"))
            self.assertEqual(summary["status"], "PASS")
        for qx in ("Q1", "Q2", "Q3"):
            metrics = json.loads((ROOT / "results" / qx / "experiments" / "round1" /
                                  "metrics" / "metrics.json").read_text(encoding="utf-8"))
            chosen = metrics.get("methods", {}).get("m3") if qx != "Q1" else metrics
            audit = chosen.get("audit", chosen)
            self.assertEqual(audit.get("simultaneous_charge_discharge_count", 0), 0)
            self.assertLessEqual(audit.get("max_balance_violation_kwh", audit.get("balance_violation_kwh", 0)), 1e-6)
        q4 = json.loads((ROOT / "results" / "Q4" / "experiments" / "round1" /
                         "metrics" / "metrics.json").read_text(encoding="utf-8"))
        self.assertEqual(q4["q4_2"]["m3"]["future_price_leakage_count"], 0)
        self.assertEqual(q4["q4_3"]["m3"]["future_price_leakage_count"], 0)
        self.assertGreater(q4["q4_2"]["m2"]["future_price_leakage_count"], 0)

    def test_core_figures_are_nonempty(self):
        paths = (
            ROOT / "results/Q1/experiments/round1/figures/q1_dispatch.png",
            ROOT / "results/Q2/experiments/round1/figures/q2_method_comparison.png",
            ROOT / "results/Q3/experiments/round1/figures/q3_method_comparison.png",
            ROOT / "results/Q4/experiments/round1/figures/q4_price_policy_comparison.png",
        )
        for path in paths:
            self.assertGreater(path.stat().st_size, 10_000, str(path))

    def test_round2_uses_selected_q3_reserve_and_preserves_constraints(self):
        q3 = json.loads((ROOT / "results/Q3/experiments/round2/metrics/metrics.json").read_text(encoding="utf-8"))
        q4 = json.loads((ROOT / "results/Q4/experiments/round2/metrics/metrics.json").read_text(encoding="utf-8"))
        self.assertEqual(q3["round"], 2)
        self.assertAlmostEqual(q3["reserve_alpha"], 0.99)
        self.assertAlmostEqual(q4["reserve_alpha_q4_3"], 0.99)
        self.assertEqual(q3["methods"]["m3"]["simultaneous_charge_discharge_count"], 0)
        self.assertLessEqual(q3["methods"]["m3"]["max_balance_violation_kwh"], 1e-6)
        self.assertEqual(q4["q4_3"]["m3"]["future_price_leakage_count"], 0)
        for name in ("result3.xlsx", "result4-3.xlsx"):
            result = load_workbook(ROOT / "output" / name, read_only=True, data_only=True)
            template = load_workbook(ROOT / "data/templates" / name, read_only=True)
            self.assertEqual(result.sheetnames, template.sheetnames)


if __name__ == "__main__":
    unittest.main()
