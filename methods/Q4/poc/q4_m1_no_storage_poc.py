import json, time
from pathlib import Path
import numpy as np
from openpyxl import load_workbook

start=time.perf_counter(); root=Path(__file__).resolve().parents[3]
wb=load_workbook(root/"data/raw/附件2.xlsx",read_only=True,data_only=True)
load=np.array(next(wb["小区负载"].iter_rows(min_row=2,max_row=2,min_col=2,values_only=True)),float)
pv=np.array(next(wb["光伏发电实际功率"].iter_rows(min_row=2,max_row=2,min_col=2,values_only=True)),float)
pws=load_workbook(root/"data/raw/附件4.xlsx",read_only=True,data_only=True).active
price=np.array(next(pws.iter_rows(min_row=2,max_row=2,min_col=2,values_only=True)),float)
net=(load-pv)/6; grid=np.maximum(net,0)
out={"candidate":"Q4-M1","cost_yuan":float(price@grid),"terminal_error_kwh":0.0,
     "balance_violation_kwh":float(np.maximum(net-grid,0).max()),"runtime_s":time.perf_counter()-start}
print(json.dumps(out))
