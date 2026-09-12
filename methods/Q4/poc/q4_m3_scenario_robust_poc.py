import json, time
from pathlib import Path
import numpy as np
from openpyxl import load_workbook
from scipy.optimize import linprog

start=time.perf_counter(); root=Path(__file__).resolve().parents[3]; S=3; n=144; eta=.9
wb=load_workbook(root/"data/raw/附件2.xlsx",read_only=True,data_only=True)
load=np.array(list(wb["小区负载"].iter_rows(min_row=2,max_row=4,min_col=2,values_only=True)),float)
pv=np.array(list(wb["光伏发电实际功率"].iter_rows(min_row=2,max_row=4,min_col=2,values_only=True)),float); net=(load-pv)/6
pws=load_workbook(root/"data/raw/附件4.xlsx",read_only=True,data_only=True).active
p=np.array(list(pws.iter_rows(min_row=2,max_row=4,min_col=2,values_only=True)),float); N=S*n+2*n+1; rows=[]; rhs=[]
for s in range(S):
    for t in range(n):
        row=np.zeros(N); row[s*n+t]=-1; row[S*n+t]=1; row[S*n+n+t]=-1; rows.append(row); rhs.append(-net[s,t])
    row=np.zeros(N); row[s*n:(s+1)*n]=p[s]; row[-1]=-1; rows.append(row); rhs.append(0)
L=np.tril(np.ones((n,n))); Z=np.zeros((n,S*n)); rows+=list(np.hstack([Z,eta*L,-L/eta,np.zeros((n,1))])); rhs+=list(np.full(n,4800))
rows+=list(np.hstack([Z,-eta*L,L/eta,np.zeros((n,1))])); rhs+=list(np.full(n,4800))
eq=np.r_[np.zeros(S*n),np.full(n,eta),np.full(n,-1/eta),0][None,:]
res=linprog(np.r_[np.zeros(N-1),1],A_ub=np.array(rows),b_ub=np.array(rhs),A_eq=eq,b_eq=[0],bounds=[(0,None)]*(S*n)+[(0,5000/6)]*(2*n)+[(0,None)],method="highs")
g=res.x[:S*n].reshape(S,n); c=res.x[S*n:S*n+n]; d=res.x[S*n+n:-1]; costs=(p*g).sum(1)
out={"candidate":"Q4-M3","worst_cost_yuan":float(costs.max()),"mean_cost_yuan":float(costs.mean()),"terminal_error_kwh":float(abs(eta*c.sum()-d.sum()/eta)),"balance_violation_kwh":float(np.maximum(net-(g+d-c),0).max()),"runtime_s":time.perf_counter()-start}
print(json.dumps(out))
