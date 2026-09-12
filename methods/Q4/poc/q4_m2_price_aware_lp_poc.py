import json, time
from pathlib import Path
import numpy as np
from openpyxl import load_workbook
from scipy.optimize import linprog

start=time.perf_counter(); root=Path(__file__).resolve().parents[3]
wb=load_workbook(root/"data/raw/附件2.xlsx",read_only=True,data_only=True)
load=np.array(next(wb["小区负载"].iter_rows(min_row=2,max_row=2,min_col=2,values_only=True)),float)
pv=np.array(next(wb["光伏发电实际功率"].iter_rows(min_row=2,max_row=2,min_col=2,values_only=True)),float)
pws=load_workbook(root/"data/raw/附件4.xlsx",read_only=True,data_only=True).active
p=np.array(next(pws.iter_rows(min_row=2,max_row=2,min_col=2,values_only=True)),float); net=(load-pv)/6; n=144
L=np.tril(np.ones((n,n))); Z=np.zeros((n,n)); eta=.9
A=np.vstack([np.hstack([-np.eye(n),np.eye(n),-np.eye(n)]),np.hstack([Z,eta*L,-L/eta]),np.hstack([Z,-eta*L,L/eta])])
b=np.r_[-net,np.full(n,4800),np.full(n,4800)]; eq=np.r_[np.zeros(n),np.full(n,eta),np.full(n,-1/eta)][None,:]
res=linprog(np.r_[p,np.zeros(2*n)],A_ub=A,b_ub=b,A_eq=eq,b_eq=[0],bounds=[(0,None)]*n+[(0,5000/6)]*(2*n),method="highs")
g,c,d=res.x[:n],res.x[n:2*n],res.x[2*n:]; base=float(p@np.maximum(net,0)); cost=float(p@g)
out={"candidate":"Q4-M2","cost_yuan":cost,"savings_pct":100*(base-cost)/base,
     "terminal_error_kwh":float(abs(.9*c.sum()-d.sum()/.9)),"balance_violation_kwh":float(np.maximum(net-(g+d-c),0).max()),"runtime_s":time.perf_counter()-start}
print(json.dumps(out))
