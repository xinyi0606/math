import json, time
from pathlib import Path
import numpy as np
from openpyxl import load_workbook

start=time.perf_counter(); root=Path(__file__).resolve().parents[3]
fws=load_workbook(root/"data/raw/附件3.xlsx",read_only=True,data_only=True).active
fc=np.array(list(fws.iter_rows(min_row=2,min_col=3,values_only=True)),float).reshape(365,4,24)
awb=load_workbook(root/"data/raw/附件2.xlsx",read_only=True,data_only=True)
actual=np.array(list(awb["光伏发电实际功率"].iter_rows(min_row=2,min_col=2,values_only=True)),float)[:,5::6]
latest=np.hstack([fc[:,r,:6] for r in range(4)])
rho=float(np.quantile((latest[:21]-actual[:21]).ravel(),.9))
safe=np.maximum(latest[21:35]-rho,0); truth=actual[21:35]
out={"candidate":"Q3-M3","reserve_kw":rho,"coverage":float(np.mean(truth>=safe)),
     "emergency_proxy_kwh":float(np.maximum(safe-truth,0).sum()),"evaluation_window":"2025-01-22..2025-02-04","runtime_s":time.perf_counter()-start}
print(json.dumps(out))
