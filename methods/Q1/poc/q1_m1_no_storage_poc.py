import json
import time
from pathlib import Path
import numpy as np
from openpyxl import load_workbook

start = time.perf_counter()
root = Path(__file__).resolve().parents[3]
ws = load_workbook(root / "data/raw/附件1.xlsx", read_only=True, data_only=True).active
rows = list(ws.iter_rows(min_row=2, values_only=True))
price = np.array([r[1] for r in rows], float)
net = (np.array([r[2] for r in rows]) - np.array([r[3] for r in rows])) / 6
grid = np.maximum(net, 0)
out = {"candidate": "Q1-M1", "cost_yuan": float(price @ grid),
       "terminal_error_kwh": 0.0, "balance_violation_kwh": float(np.maximum(net-grid, 0).max()),
       "runtime_s": time.perf_counter() - start}
print(json.dumps(out))
