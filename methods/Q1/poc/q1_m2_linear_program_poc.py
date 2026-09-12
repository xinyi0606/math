import json, time
from pathlib import Path
import numpy as np
from openpyxl import load_workbook
from scipy.optimize import linprog

start=time.perf_counter(); root=Path(__file__).resolve().parents[3]
ws=load_workbook(root/"data/raw/附件1.xlsx",read_only=True,data_only=True).active
rows=list(ws.iter_rows(min_row=2,values_only=True)); n=len(rows)
p=np.array([r[1] for r in rows]); net=(np.array([r[2] for r in rows])-np.array([r[3] for r in rows]))/6
L=np.tril(np.ones((n,n))); Z=np.zeros((n,n)); eta=.9
A=np.vstack([np.hstack([-np.eye(n),np.eye(n),-np.eye(n)]),
             np.hstack([Z,eta*L,-L/eta]),np.hstack([Z,-eta*L,L/eta])])
b=np.r_[-net,np.full(n,4800),np.full(n,4800)]
eq=np.r_[np.zeros(n),np.full(n,eta),np.full(n,-1/eta)][None,:]
res=linprog(np.r_[p,np.zeros(2*n)],A_ub=A,b_ub=b,A_eq=eq,b_eq=[0],
            bounds=[(0,None)]*n+[(0,5000/6)]*(2*n),method="highs")
g,c,d=res.x[:n],res.x[n:2*n],res.x[2*n:]
out={"candidate":"Q1-M2","cost_yuan":float(p@g),"terminal_error_kwh":float(abs(eta*c.sum()-d.sum()/eta)),
     "balance_violation_kwh":float(np.maximum(net-(g+d-c),0).max()),"runtime_s":time.perf_counter()-start}
print(json.dumps(out))
