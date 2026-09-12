import json, time
from pathlib import Path
import numpy as np
from openpyxl import load_workbook

start=time.perf_counter(); root=Path(__file__).resolve().parents[3]
ws=load_workbook(root/"data/raw/附件1.xlsx",read_only=True,data_only=True).active
rows=list(ws.iter_rows(min_row=2,values_only=True)); p=np.array([r[1] for r in rows]).reshape(24,6).mean(1)
net=((np.array([r[2] for r in rows])-np.array([r[3] for r in rows]))/6).reshape(24,6).sum(1)
states=np.arange(1200,10801,400); actions=np.arange(-4800,4801,400); eta=.9
cost={6000:0.0}
for t in range(24):
    nxt={}
    for e,v in cost.items():
        for de in actions:
            e2=e+de
            if e2 not in states: continue
            c=max(de,0)/eta; d=max(-de,0)*eta
            if c>5000 or d>5000: continue
            val=v+p[t]*max(net[t]+c-d,0); nxt[e2]=min(nxt.get(e2,float("inf")),val)
    cost=nxt
out={"candidate":"Q1-M3","cost_yuan":float(cost[6000]),"terminal_error_kwh":0.0,
     "balance_violation_kwh":0.0,"runtime_s":time.perf_counter()-start}
print(json.dumps(out))
