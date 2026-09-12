import json, time
from pathlib import Path
import numpy as np
from openpyxl import load_workbook

start=time.perf_counter(); root=Path(__file__).resolve().parents[3]
wb=load_workbook(root/"data/raw/附件2.xlsx",read_only=True,data_only=True)
load=np.array(list(wb["小区负载"].iter_rows(min_row=2,min_col=2,values_only=True)),float)
pv=np.array(list(wb["光伏发电实际功率"].iter_rows(min_row=2,min_col=2,values_only=True)),float)
price=np.array([r[0] for r in load_workbook(root/"data/raw/附件1.xlsx",read_only=True,data_only=True)["Sheet1"].iter_rows(min_row=2,min_col=2,max_col=2,values_only=True)],float)
net=load-pv; days=range(21,28)
actual=np.vstack([net[d] for d in days]); forecast=np.vstack([net[d-7:d].mean(0) for d in days])
err=actual-forecast; short=np.maximum(err,0)
out={"candidate":"Q2-M2","rmse_kw":float(np.sqrt(np.mean(err**2))),
     "mae_kw":float(np.mean(abs(err))),"underprediction_kwh":float(short.sum()/6),
     "emergency_cost_proxy_yuan":float((short*price).sum()/6*5),"underprediction_p95_kw":float(np.quantile(short,0.95)),
     "evaluation_window":"2025-01-22..2025-01-28","runtime_s":time.perf_counter()-start}
print(json.dumps(out))
